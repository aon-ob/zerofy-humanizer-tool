import io
import re
import docx
import pypdf

MAX_CHUNK_CHARACTERS = 13500

# Common AI-correlated tokens and filler words prioritized during Stage 1 pruning
AI_MARKER_WORDS = {
    "moreover", "furthermore", "imperative", "comprehensive", "meticulously",
    "facilitating", "facilitate", "paradigm", "delve", "testament", "pivotal",
    "underscore", "underscores", "realm", "foster", "fosters", "noteworthy",
    "crucial", "essential", "vital", "significant", "significantly", "effectively",
    "ultimately", "consequently", "additionally", "specifically", "particularly",
    "thoroughly", "seamless", "seamlessly", "orchestrate", "leverage", "utilize",
    "important", "described", "certain", "designed", "therefore", "however"
}

# Grammar words we avoid deleting first unless no other token unflags the segment
LOW_PRIORITY_GRAMMAR_WORDS = {
    "a", "an", "the", "is", "are", "was", "were", "be", "to", "of", "in", "on",
    "for", "with", "as", "by", "at", "from", "or", "and", "if", "it", "not"
}


def split_text_into_chunks(text: str, max_size: int = MAX_CHUNK_CHARACTERS) -> list[str]:
    # Splits large text into chunks <= max_size preserving paragraph and sentence boundaries
    if len(text) <= max_size:
        return [text.strip()] if text.strip() else []

    paragraphs = text.split("\n\n")
    chunks = []
    current_chunk = []
    current_length = 0

    for paragraph in paragraphs:
        p_len = len(paragraph) + 2
        if p_len > max_size:
            sentences = split_into_sentences(paragraph)
            for sentence in sentences:
                s_len = len(sentence) + 1
                if current_length + s_len > max_size and current_chunk:
                    chunks.append("\n\n".join(current_chunk))
                    current_chunk = []
                    current_length = 0
                current_chunk.append(sentence)
                current_length += s_len
        else:
            if current_length + p_len > max_size and current_chunk:
                chunks.append("\n\n".join(current_chunk))
                current_chunk = []
                current_length = 0
            current_chunk.append(paragraph)
            current_length += p_len

    if current_chunk:
        chunks.append("\n\n".join(current_chunk))

    return [c.strip() for c in chunks if c.strip()]


def split_into_sentences(text: str) -> list[str]:
    # Splits paragraph text into sentences while respecting common abbreviations
    if not text.strip():
        return []
    protected = re.sub(
        r"\b(Mr|Mrs|Ms|Dr|Prof|Sr|Jr|vs|etc|e\.g|i\.e)\.",
        r"\1<PRD>",
        text,
        flags=re.IGNORECASE,
    )
    raw_sentences = re.split(r"(?<=[.!?])\s+", protected)
    return [s.replace("<PRD>", ".").strip() for s in raw_sentences if s.strip()]


def segment_in_text(text: str, target: str) -> bool:
    # Checks if target exists in text, including across newlines or multiple spaces
    if not target or not text:
        return False
    if target in text:
        return True
    words = target.split()
    if not words:
        return False
    flex_pattern = r"\s+".join(re.escape(w) for w in words)
    return bool(re.search(flex_pattern, text))


def extract_enclosing_sentence(text: str, target: str) -> str:
    # Finds the exact enclosing sentence for a target segment using word-boundary matching
    if not target or not text:
        return target

    clean_target = target.strip()
    if not clean_target:
        return target

    paragraphs = text.split("\n\n")
    use_word_boundary = len(clean_target) < 18 and bool(re.match(r"^[\w\s]+$", clean_target))
    if use_word_boundary:
        pattern = re.compile(r"\b" + r"\s+".join(re.escape(w) for w in clean_target.split()) + r"\b")
    else:
        pattern = re.compile(r"\s+".join(re.escape(w) for w in clean_target.split()))

    for p in paragraphs:
        for s in split_into_sentences(p):
            if clean_target == s or pattern.search(s):
                return s
        if pattern.search(p):
            return p.strip()

    return clean_target


def extract_enclosing_paragraph(text: str, target_sentence: str) -> str:
    # Returns the full paragraph containing target_sentence for LLM context
    paragraphs = text.split("\n\n")
    for p in paragraphs:
        if segment_in_text(p, target_sentence):
            return p.strip()
    return text[:1500].strip()


def read_file_text(file_bytes: bytes, filename: str) -> str:
    # Extracts plain text from .txt, .md, .docx, and .pdf files
    ext = filename.lower().split(".")[-1]
    if ext in ("txt", "md", "csv"):
        text = file_bytes.decode("utf-8", errors="ignore")
    elif ext == "docx":
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    elif ext == "pdf":
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join([pt.strip() for pt in pages_text if pt.strip()])
    else:
        raise ValueError("Unsupported format. Only .txt, .md, .docx, and .pdf files are supported.")

    return text.strip()


def create_docx_bytes(text: str) -> bytes:
    # Creates a downloadable .docx file in memory from clean text
    doc = docx.Document()
    for para in text.split("\n\n"):
        cleaned = para.strip()
        if cleaned:
            doc.add_paragraph(cleaned)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _replace_once_bounded(text: str, target: str, replacement: str) -> str:
    if not target:
        return text
    if target in text and len(target) >= 15:
        return text.replace(target, replacement, 1)

    prefix = r"\b" if re.match(r"^\w", target) else ""
    suffix = r"\b" if re.search(r"\w$", target) else ""
    pattern = prefix + re.escape(target) + suffix
    try:
        new_text, count = re.subn(pattern, lambda _: replacement, text, count=1)
        if count > 0:
            return new_text
    except Exception:
        pass

    words = target.split()
    if words:
        flex_pattern = prefix + r"\s+".join(re.escape(w) for w in words) + suffix
        try:
            new_text, count = re.subn(flex_pattern, lambda _: replacement, text, count=1)
            if count > 0:
                return new_text
        except Exception:
            pass

    return text.replace(target, replacement, 1)


def replace_segment_in_text(full_text: str, target: str, replacement: str, enclosing_sentence: str = "") -> str:
    # Replaces ONLY the single target occurrence, scoped to enclosing_sentence when provided
    if not target:
        return full_text

    if enclosing_sentence and segment_in_text(full_text, enclosing_sentence) and target != enclosing_sentence:
        updated_sentence = _replace_once_bounded(enclosing_sentence, target, replacement)
        updated_sentence = re.sub(r"\s+([,.;:!?])", r"\1", updated_sentence)
        updated_sentence = re.sub(r" {2,}", " ", updated_sentence).strip()
        return _replace_once_bounded(full_text, enclosing_sentence, updated_sentence)

    res = _replace_once_bounded(full_text, target, replacement)
    res = re.sub(r"\s+([,.;:!?])", r"\1", res)
    return re.sub(r" {2,}", " ", res).strip()


def score_ablation_candidate(words: list[str], start_idx: int, k: int) -> float:
    # Prioritizes AI marker words, adverbs, and interior modifiers over sentence starters
    chunk = words[start_idx : start_idx + k]
    cleaned_words = [re.sub(r"[^\w]", "", w).lower() for w in chunk]
    score = 0.0

    for cw in cleaned_words:
        if cw in AI_MARKER_WORDS:
            score += 50.0
        elif cw.endswith("ly") and len(cw) > 4:
            score += 30.0
        elif len(cw) >= 7:
            score += 15.0
        elif cw in LOW_PRIORITY_GRAMMAR_WORDS:
            score -= 20.0

    if start_idx == 0:
        score -= 25.0

    raw_joined = " ".join(chunk)
    if "," in raw_joined:
        score += 12.0

    return score


def build_prioritized_ablation_candidates(current_text: str, target_segment: str, k: int) -> list[dict]:
    # Builds and sorts candidate removals of length k by AI trigger likelihood
    words = target_segment.split()
    if k > len(words):
        return []

    raw_candidates = []
    for i in range(len(words) - k + 1):
        removed_chunk = words[i : i + k]
        removed_str = " ".join(removed_chunk)
        modified_words = words[:i] + words[i + k :]
        modified_segment = " ".join(modified_words)
        modified_segment = re.sub(r"\s+([,.;:!?])", r"\1", modified_segment).strip()

        candidate_text = replace_segment_in_text(current_text, target_segment, modified_segment)
        if candidate_text == current_text:
            continue

        priority = score_ablation_candidate(words, i, k)

        raw_candidates.append({
            "priority": priority,
            "index": i,
            "removed_str": removed_str,
            "modified_segment": modified_segment,
            "candidate_text": candidate_text,
            "source_sentence": target_segment,
        })

    raw_candidates.sort(key=lambda item: (-item["priority"], item["index"]))
    return raw_candidates


def is_segment_highlighted(segment: str, highlights: list[str]) -> bool:
    # Bidirectional check used in Stage 2: returns True if segment overlaps with any ZeroGPT highlight
    if not segment or not segment.strip():
        return False
    norm_seg = re.sub(r"\s+", " ", segment).strip().lower()
    if len(norm_seg) < 4:
        return False
    for h in highlights:
        norm_h = re.sub(r"\s+", " ", h).strip().lower()
        if not norm_h:
            continue
        if norm_seg in norm_h or norm_h in norm_seg:
            return True
    return False


def generate_highlighted_html(original_text: str, modified_phrases: list[dict]) -> str:
    # Generates safe interactive diff HTML using placeholders scoped to source_sentence
    working_text = original_text
    placeholders = {}

    for idx, item in enumerate(modified_phrases):
        phrase = item.get("original", "")
        replacement_text = item.get("replacement", "")
        status = item.get("status", "ablated")
        source_sentence = item.get("source_sentence", "")

        if not phrase or not phrase.strip():
            continue

        token_key = f"__ZEROFY_DIFF_TOKEN_{idx}__"
        placeholders[token_key] = {
            "idx": idx,
            "original": phrase,
            "replacement": replacement_text,
            "status": status,
        }

        if source_sentence and segment_in_text(working_text, source_sentence) and segment_in_text(source_sentence, phrase) and phrase != source_sentence:
            updated_sent = _replace_once_bounded(source_sentence, phrase, token_key)
            working_text = _replace_once_bounded(working_text, source_sentence, updated_sent)
        else:
            working_text = _replace_once_bounded(working_text, phrase, token_key)

    escaped_html = (
        working_text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    for token_key, meta in placeholders.items():
        idx = meta["idx"]
        orig_esc = (
            meta["original"]
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )
        repl_esc = (
            meta["replacement"]
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )
        status = meta["status"]

        if status == "replaced" and repl_esc:
            mark_html = (
                f'<mark data-diff-idx="{idx}" onclick="openDiffInspector({idx})" '
                f'class="diff-mark bg-emerald-950/85 text-emerald-300 px-1.5 py-0.5 rounded border border-emerald-600/70 font-medium cursor-pointer hover:bg-emerald-900 transition-all" '
                f'title="Original: {orig_esc} (Click to inspect or edit)">{repl_esc}</mark>'
            )
        elif status == "manual":
            mark_html = (
                f'<mark data-diff-idx="{idx}" onclick="openDiffInspector({idx})" '
                f'class="diff-mark bg-amber-950/85 text-amber-300 px-1.5 py-0.5 rounded border border-amber-500/80 font-medium cursor-pointer hover:bg-amber-900 transition-all" '
                f'title="Manual Review Needed - Click to edit or re-roll with Ollama">{orig_esc}</mark>'
            )
        else:
            mark_html = (
                f'<mark data-diff-idx="{idx}" onclick="openDiffInspector({idx})" '
                f'class="diff-mark bg-rose-950/70 text-rose-300 px-1.5 py-0.5 rounded border border-rose-700/60 line-through decoration-rose-400 mx-0.5 cursor-pointer hover:bg-rose-900/80 transition-all" '
                f'title="Pruned AI trigger token (Click to restore or edit)">{orig_esc}</mark>'
            )

        escaped_html = escaped_html.replace(token_key, mark_html)

    paragraphs = escaped_html.split("\n\n")
    formatted_paragraphs = []
    for p in paragraphs:
        if p.strip():
            p_clean = p.replace("\n", "<br>")
            formatted_paragraphs.append(f"<p class='mb-4 leading-relaxed'>{p_clean}</p>")

    return "".join(formatted_paragraphs)